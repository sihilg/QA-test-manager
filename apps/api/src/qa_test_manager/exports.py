from __future__ import annotations

import re
import unicodedata
from collections.abc import Sequence
from io import BytesIO
from typing import Annotated, Literal
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.flowables import Flowable
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from qa_test_manager.auth import AuthContext, get_auth_context
from qa_test_manager.database import get_session
from qa_test_manager.models import Project, TestCase, TestPriority, TestResult, TestType
from qa_test_manager.test_cases import _accessible_project, _case_or_404

router = APIRouter(tags=["exports"])

ExportFormat = Literal["docx", "xlsx", "pdf"]
MAX_EXPORT_CASES = 1000
MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


@router.get("/test-cases/{case_id}/export/{export_format}")
def export_test_case(
    case_id: int,
    export_format: ExportFormat,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    database: Annotated[Session, Depends(get_session)],
) -> Response:
    test_case = _case_or_404(database, case_id)
    project = _accessible_project(database, context, test_case.project_id)
    test_case = _load_case(database, case_id)
    filename = f"{test_case.case_number}-{_slug(test_case.title)}.{export_format}"
    return _export_response([test_case], project, export_format, filename)


@router.get("/projects/{project_id}/test-cases/export/{export_format}")
def export_project_test_cases(
    project_id: int,
    export_format: ExportFormat,
    context: Annotated[AuthContext, Depends(get_auth_context)],
    database: Annotated[Session, Depends(get_session)],
    result: TestResult | None = None,
    priority: TestPriority | None = None,
    test_type: TestType | None = None,
    limit: Annotated[int, Query(ge=1, le=MAX_EXPORT_CASES)] = MAX_EXPORT_CASES,
) -> Response:
    project = _accessible_project(database, context, project_id)
    statement = (
        select(TestCase)
        .options(selectinload(TestCase.steps))
        .where(TestCase.project_id == project_id, TestCase.archived_at.is_(None))
    )
    if result is not None:
        statement = statement.where(TestCase.final_result == result)
    if priority is not None:
        statement = statement.where(TestCase.priority == priority)
    if test_type is not None:
        statement = statement.where(TestCase.test_type == test_type)
    cases = list(database.scalars(statement.order_by(TestCase.execution_order).limit(limit + 1)))
    if len(cases) > limit:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"A exportação está limitada a {limit} casos.",
        )
    filename = f"{_slug(project.name)}-casos.{export_format}"
    return _export_response(cases, project, export_format, filename)


def _load_case(database: Session, case_id: int) -> TestCase:
    test_case = database.scalar(
        select(TestCase).options(selectinload(TestCase.steps)).where(TestCase.id == case_id)
    )
    if test_case is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Caso não encontrado.")
    return test_case


def _export_response(
    cases: Sequence[TestCase], project: Project, export_format: ExportFormat, filename: str
) -> Response:
    generators = {"docx": _build_docx, "xlsx": _build_xlsx, "pdf": _build_pdf}
    content = generators[export_format](cases, project)
    encoded = quote(filename)
    disposition = (
        f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded}"
    )
    return Response(
        content=content,
        media_type=MIME_TYPES[export_format],
        headers={
            "Content-Disposition": disposition,
            "X-Content-Type-Options": "nosniff",
        },
    )


def _build_docx(cases: Sequence[TestCase], project: Project) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(22)
    document.add_heading(f"Casos de teste de {project.name}", 0)
    document.add_paragraph(
        f"Exportação local com {len(cases)} caso(s) ativo(s), ordenados para execução."
    )
    for index, test_case in enumerate(cases):
        if index:
            document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading(f"{test_case.case_number} {test_case.title}", level=1)
        metadata = [
            ("Ordem", str(test_case.execution_order)),
            ("Prioridade", test_case.priority.value),
            ("Tipo", test_case.test_type.value),
            ("Resultado", test_case.final_result.value),
            ("Requisito", test_case.requirement),
            ("Browser", test_case.browser or "Não informado"),
            ("Data de execução", str(test_case.execution_date or "Não informada")),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in metadata:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[0].paragraphs[0].runs[0].bold = True
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_shading(
                    cell,
                    "F3F4F6"
                    if label in {"Ordem", "Tipo", "Requisito", "Data de execução"}
                    else "FFFFFF",
                )
        for heading, value in [
            ("Descrição", test_case.description),
            ("Dados de teste", test_case.test_data),
            ("Pré-condições", test_case.preconditions),
            ("Resultado esperado", test_case.expected_result),
        ]:
            document.add_heading(heading, level=2)
            document.add_paragraph(value)
        document.add_heading("Passos", level=2)
        steps = document.add_table(rows=1, cols=3)
        steps.style = "Table Grid"
        headers = steps.rows[0].cells
        for cell, text in zip(headers, ["#", "Ação", "Resultado esperado"], strict=True):
            cell.text = text
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _set_cell_shading(cell, "1F4E78")
        for step in sorted(test_case.steps, key=lambda item: item.position):
            cells = steps.add_row().cells
            cells[0].text = str(step.position)
            cells[1].text = step.action
            cells[2].text = step.expected_result or ""
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _build_xlsx(cases: Sequence[TestCase], project: Project) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "Casos"
    sheet.sheet_view.showGridLines = False
    headers = [
        "Número",
        "Ordem",
        "Título",
        "Prioridade",
        "Tipo",
        "Resultado",
        "Requisito",
        "Browser",
        "Data de execução",
        "Dados de teste",
        "Descrição",
        "Pré-condições",
        "Resultado esperado",
        "Passos",
    ]
    sheet.append([f"Casos de teste de {project.name}"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="000000")
    sheet.row_dimensions[1].height = 24
    sheet.append(headers)
    sheet.row_dimensions[2].height = 30
    for cell in sheet[2]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for test_case in cases:
        steps = "\n".join(
            f"{step.position}. {step.action} | {step.expected_result or ''}"
            for step in sorted(test_case.steps, key=lambda item: item.position)
        )
        sheet.append(
            [
                _excel_safe(test_case.case_number),
                test_case.execution_order,
                _excel_safe(test_case.title),
                test_case.priority.value,
                test_case.test_type.value,
                test_case.final_result.value,
                _excel_safe(test_case.requirement),
                _excel_safe(test_case.browser or ""),
                test_case.execution_date,
                _excel_safe(test_case.test_data),
                _excel_safe(test_case.description),
                _excel_safe(test_case.preconditions),
                _excel_safe(test_case.expected_result),
                _excel_safe(steps),
            ]
        )
    widths = [12, 9, 30, 13, 24, 16, 20, 16, 18, 28, 32, 28, 32, 48]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_index, row in enumerate(sheet.iter_rows(min_row=3), start=3):
        sheet.row_dimensions[row_index].height = 48
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:N{max(sheet.max_row, 2)}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_pdf(cases: Sequence[TestCase], project: Project) -> bytes:
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=letter,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Casos de teste de {project.name}",
        author="QA Test Manager",
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "Helvetica-Bold"
    styles["Title"].fontSize = 20
    styles["Title"].textColor = colors.black
    styles["Title"].alignment = TA_CENTER
    body = ParagraphStyle(
        "ExportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13
    )
    heading = ParagraphStyle(
        "ExportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=colors.black,
        spaceBefore=8,
        spaceAfter=4,
    )
    table_header = ParagraphStyle(
        "ExportTableHeader",
        parent=body,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    story: list[Flowable] = [
        Paragraph(_pdf_text(f"Casos de teste de {project.name}"), styles["Title"]),
        Spacer(1, 5 * mm),
        Paragraph(_pdf_text(f"Exportação local com {len(cases)} caso(s) ativo(s)."), body),
        Spacer(1, 6 * mm),
    ]
    for index, test_case in enumerate(cases):
        if index:
            story.append(PageBreak())
        story.append(
            Paragraph(_pdf_text(f"{test_case.case_number} {test_case.title}"), styles["Heading1"])
        )
        data = [
            ["Ordem", str(test_case.execution_order), "Prioridade", test_case.priority.value],
            ["Tipo", test_case.test_type.value, "Resultado", test_case.final_result.value],
            ["Requisito", test_case.requirement, "Browser", test_case.browser or "Não informado"],
            [
                "Execução",
                str(test_case.execution_date or "Não informada"),
                "Origem",
                test_case.origin.value,
            ],
        ]
        metadata = Table(
            [[Paragraph(_pdf_text(str(value)), body) for value in row] for row in data],
            colWidths=[24 * mm, 55 * mm, 24 * mm, 55 * mm],
        )
        metadata.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2F8")),
                    ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF2F8")),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.extend([metadata, Spacer(1, 4 * mm)])
        for label, value in [
            ("Descrição", test_case.description),
            ("Dados de teste", test_case.test_data),
            ("Pré-condições", test_case.preconditions),
            ("Resultado esperado", test_case.expected_result),
        ]:
            story.extend([Paragraph(label, heading), Paragraph(_pdf_text(value), body)])
        story.append(Paragraph("Passos", heading))
        rows = [
            [
                Paragraph("#", table_header),
                Paragraph("Ação", table_header),
                Paragraph("Resultado esperado", table_header),
            ]
        ]
        rows.extend(
            [
                [
                    Paragraph(str(step.position), body),
                    Paragraph(_pdf_text(step.action), body),
                    Paragraph(_pdf_text(step.expected_result or ""), body),
                ]
                for step in sorted(test_case.steps, key=lambda item: item.position)
            ]
        )
        steps_table = Table(rows, colWidths=[10 * mm, 74 * mm, 74 * mm], repeatRows=1)
        steps_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(steps_table)
    document.build(story)
    return stream.getvalue()


def _set_cell_shading(cell: object, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _excel_safe(value: str) -> str:
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
        return "'" + value
    return value


def _pdf_text(value: str) -> str:
    return (
        value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
    )


def _slug(value: str) -> str:
    normalized = (
        unicodedata.normalize("NFKD", value.lower()).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return (slug or "exportacao")[:80]
