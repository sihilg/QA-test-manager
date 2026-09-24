from collections.abc import Iterator
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from qa_test_manager.database import get_session
from qa_test_manager.main import app
from qa_test_manager.models import Base, Project, User, UserRole
from qa_test_manager.security import hash_password


@pytest.fixture
def export_context() -> Iterator[tuple[TestClient, int]]:
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as database:
        admin = User(
            name="Admin",
            email="admin@example.com",
            password_hash=hash_password("admin-password-123"),
            role=UserRole.ADMIN,
        )
        project = Project(name="Portal Português")
        database.add_all([admin, project])
        database.commit()
        project_id = project.id

    def override_session() -> Iterator[Session]:
        with factory() as database:
            yield database

    app.dependency_overrides[get_session] = override_session
    with TestClient(app) as client:
        yield client, project_id
    app.dependency_overrides.clear()


def _login(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/auth/login", json={"email": "admin@example.com", "password": "admin-password-123"}
    )
    return {"X-CSRF-Token": str(response.json()["csrf_token"])}


def _case_payload(order: int = 1, title: str = "Autenticação válida") -> dict[str, object]:
    return {
        "execution_order": order,
        "priority": "HIGH" if order == 1 else "LOW",
        "requirement": "RF-001",
        "test_type": "FUNCTIONAL_POSITIVE",
        "title": title,
        "test_data": "utilizador português",
        "description": "Validar autenticação com acentuação",
        "preconditions": "Utilizador ativo",
        "expected_result": "A sessão é iniciada",
        "final_result": "NOT_EXECUTED",
        "origin": "MANUAL",
        "steps": [
            {"position": 1, "action": "Submeter credenciais", "expected_result": "Sessão iniciada"}
        ],
    }


def _create_case(client: TestClient, project_id: int, **changes: object) -> dict[str, object]:
    payload = _case_payload()
    payload.update(changes)
    response = client.post(
        f"/projects/{project_id}/test-cases", json=payload, headers=_login(client)
    )
    assert response.status_code == 201
    return dict(response.json())


def test_exports_one_case_as_docx_xlsx_and_pdf(
    export_context: tuple[TestClient, int],
) -> None:
    client, project_id = export_context
    test_case = _create_case(client, project_id)

    docx_response = client.get(f"/test-cases/{test_case['id']}/export/docx")
    assert docx_response.status_code == 200
    assert "CT001-autenticacao-valida.docx" in docx_response.headers["content-disposition"]
    document = Document(BytesIO(docx_response.content))
    assert "Portal Português" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert any(
        "Submeter credenciais" in cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    xlsx_response = client.get(f"/test-cases/{test_case['id']}/export/xlsx")
    workbook = load_workbook(BytesIO(xlsx_response.content), data_only=False)
    sheet = workbook["Casos"]
    assert sheet["A3"].value == "CT001"
    assert sheet["C3"].value == "Autenticação válida"

    pdf_response = client.get(f"/test-cases/{test_case['id']}/export/pdf")
    assert pdf_response.content.startswith(b"%PDF")
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(pdf_response.content)).pages
    )
    assert "Autenticação válida" in pdf_text
    assert "Submeter credenciais" in pdf_text


def test_project_export_respects_filters_and_neutralizes_excel_formulas(
    export_context: tuple[TestClient, int],
) -> None:
    client, project_id = export_context
    _create_case(client, project_id, title='=HYPERLINK("https://invalid")')
    _create_case(
        client, project_id, execution_order=2, priority="LOW", title="Caso de baixa prioridade"
    )

    response = client.get(f"/projects/{project_id}/test-cases/export/xlsx?priority=HIGH")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook["Casos"]
    assert sheet.max_row == 3
    assert sheet["C3"].value == '\'=HYPERLINK("https://invalid")'
    assert sheet["C3"].data_type == "s"


def test_export_requires_authentication_and_valid_format(
    export_context: tuple[TestClient, int],
) -> None:
    client, project_id = export_context
    assert client.get(f"/projects/{project_id}/test-cases/export/pdf").status_code == 401
    _login(client)
    assert client.get(f"/projects/{project_id}/test-cases/export/csv").status_code == 422
